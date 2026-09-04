import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("FocusNotes Project: Technical Overview & Comprehensive Report")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x71, 0xE3)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle_p.add_run("Context-Aware Conversation & Audio Intelligence Platform\nDocumentation & System Architecture")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x86, 0x86, 0x8B)

    doc.add_paragraph() # Spacing

    def add_h1(text):
        h = doc.add_heading(level=1)
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x00, 0x71, 0xE3)
        return h

    def add_h2(text):
        h = doc.add_heading(level=2)
        run = h.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
        return h

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.bold = True
        p.add_run(text)
        return p

    # 1. PROJECT IDENTITY
    add_h1("1. PROJECT IDENTITY")
    add_p("FocusNotes (also referred to as FocusNotes AI) is a context-aware conversation and audio intelligence platform engineered to automatically digest meeting transcripts and audio recordings, filter out agenda-irrelevant banter, categorize key information into actionable meeting notes (Action Items, Decisions, Problems, Ideas, Information), and generate concise, abstractive sticky note summaries.")
    add_p("Python 3.14 / 3.x, Streamlit, pandas, numpy, scikit-learn, nltk, spacy (en_core_web_sm v3.8.0), sentence-transformers, transformers (HuggingFace Hub), librosa, torch, joblib, rouge-score.", "Tech Stack & Libraries: ")

    # 2. ARCHITECTURE
    add_h1("2. ARCHITECTURE")
    add_p("The processing pipeline executes through an ordered sequence of 6 distinct stages, from raw multi-modal input to structured sticky note generation:")
    
    stages = [
        ("Stage 1: Input & Audio ASR Transcription (Optional Audio Gating)", "transcribe_audio_whisper() in src/extraction.py. When text is empty/missing and an audio file (.wav, .mp3, .m4a) is uploaded, HuggingFace Whisper (openai/whisper-tiny) transcribes the audio into text with utterance timestamp chunks."),
        ("Stage 2: Sentence Tokenization & Cleaning", "split_sentences() & clean_text() in src/preprocessing.py / FocusNotesPipeline.split_sentences() in src/pipeline.py. Cleans whitespace and segments raw text transcripts into individual sentences via nltk.sent_tokenize or regex delimiter split (?<=[.!?])\\s+."),
        ("Stage 3: Agenda Relevance Filtering (Stage 1 Classifier)", "LinearSVC pipeline loaded via models/relevance_classifier.pkl (trained via src/train_relevance.py) and RelevanceDetector in src/relevance.py using SentenceTransformer all-MiniLM-L6-v2 (cosine similarity threshold 0.20) to filter out off-topic utterances (IRRELEVANT vs RELEVANT)."),
        ("Stage 4: Information Type Classification (Stage 2 Classifier)", "InformationClassifier in src/classification.py (persisted in models/information_classifier.pkl). Classifies relevant sentences into 5 target categories: PROBLEM, IDEA, ACTION, DECISION, or INFORMATION."),
        ("Stage 5: Sticky Note Synthesis (Extractive & Abstractive)", "StickyNoteGenerator in src/sticky_notes.py and NoteSummarizer in src/summarization.py. Groups relevant sentences by category. In generative mode, passes category-tagged text to HuggingFace facebook/bart-large-cnn to produce concise bullet-point notes."),
        ("Stage 6: Extended Feature & Prosodic Extraction", "extract_named_entities(), detect_repetition(), detect_qa_pairs(), extract_speaker_turns_and_interruptions(), and extract_prosodic_features() in src/extraction.py.")
    ]
    for st_title, st_desc in stages:
        doc.add_paragraph(st_title, style='List Bullet')
        p_sub = doc.add_paragraph(st_desc)
        p_sub.paragraph_format.left_indent = Inches(0.25)

    add_h2("Optional & Conditional Gating Paths")
    add_p("Prosodic features (pitch variance and pause duration) are calculated using librosa only when audio is uploaded. When audio is missing, a fallback vector containing 0.0 values is emitted to prevent feature noise.", "Audio Gating (has_audio): ")
    add_p("Interruption detection and turn length analysis execute only when speaker metadata is present in the transcript (dict format with speaker and text keys). When unannotated plain text is passed, turn count defaults to sentence count and emits zeroed feature vectors.", "Speaker Labels Gating (has_speaker_labels): ")

    # 3. FEATURES IMPLEMENTED
    add_h1("3. FEATURES IMPLEMENTED (FULL LIST)")
    
    # Table for features
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ['Feature Name', 'Extraction Method / Library', 'Output Format', 'Gated / Conditional Status']
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shading = parse_xml(r'<w:shd {} w:fill="0071E3"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    feature_data = [
        ("Agenda Relevance Score", "TF-IDF n-grams + SentenceTransformer (all-MiniLM-L6-v2) cosine similarity in src/relevance.py", "Binary flag (0/1, RELEVANT/IRRELEVANT) & float score", "Gated on agenda text input"),
        ("Information Category Classification", "TF-IDF n-grams (1, 2) + LogisticRegression / LinearSVC in src/classification.py", "5 categories: PROBLEM, IDEA, ACTION, DECISION, INFORMATION", "Gated on relevance == 1"),
        ("Named Entity Recognition (NER)", "spaCy en_core_web_sm model in src/extraction.py", "Lists of PERSON, DATE, ORG, and GPE entity tuples", "Always computed (empty list fallback)"),
        ("Utterance Repetition & Near-Duplicates", "SentenceTransformer (all-MiniLM-L6-v2) cosine similarity matrix (threshold >= 0.85)", "is_repeated boolean flag & list of duplicate matches", "Always computed across sentences"),
        ("Q&A Pair Detection", "Regex question patterns + adjacent speaker turn matching in src/extraction.py", "List of dicts linking question and answer turns", "Always computed"),
        ("Speaker Turn & Interruption Stats", "Speaker metadata parser & timestamp overlap check in src/extraction.py", "Dict with total turns, speaker words, and interruption count", "Gated by has_speaker_labels"),
        ("Automatic Speech Recognition (ASR)", "HuggingFace openai/whisper-tiny ASR pipeline with timestamp chunking", "Transcribed text string & list of timestamp tuples", "Gated on audio upload when text is empty"),
        ("Prosodic Pitch & Pause Features", "librosa.pyin (fundamental frequency variance) & silence gap computation", "Per-utterance pitch variance (float) & pause duration (seconds)", "Gated by has_audio flag (emits 0.0 when absent)"),
        ("Abstractive Sticky Note Summarization", "HuggingFace facebook/bart-large-cnn Seq2Seq generator in src/summarization.py", "Bullet-point sticky note strings (<= 20 words)", "Gated on mode == 'generative'")
    ]

    for row_data in feature_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            row_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph() # Spacing

    # 4. MODELS USED
    add_h1("4. MODELS USED")
    models = [
        ("LinearSVC (sklearn.svm.LinearSVC)", "Stage 1 Agenda Relevance Classifier (models/relevance_classifier.pkl) and standalone benchmark model in src/train_relevance.py (C=1.0, random_state=42)."),
        ("LogisticRegression (sklearn.linear_model.LogisticRegression)", "Stage 2 Information Type Classifier (src/classification.py) for multi-class classification across relevant sentences (max_iter=1000, random_state=42)."),
        ("SentenceTransformer (all-MiniLM-L6-v2)", "Generates 384-dimensional dense semantic embeddings for calculating agenda-utterance cosine similarity in src/relevance.py and near-duplicate utterance detection in src/extraction.py."),
        ("BART Variant (facebook/bart-large-cnn)", "Sequence-to-sequence abstractive summarizer and sentence compression model in NoteSummarizer (src/summarization.py). Supported fallbacks: bart-xsum, google/pegasus-xsum, distilbart-cnn-12-6."),
        ("Whisper Variant (openai/whisper-tiny)", "Speech-to-text transcription and timestamp chunking for uploaded WAV/MP3 files in src/extraction.py."),
        ("spaCy Language Model (en_core_web_sm v3.8.0)", "Named Entity Extraction (PERSON, DATE, ORG, GPE) and structural prompt tagging in src/extraction.py and src/summarization.py.")
    ]
    for m_name, m_desc in models:
        add_p(m_desc, f"{m_name}: ")

    # 5. EVALUATION RESULTS
    add_h1("5. EVALUATION RESULTS")
    
    add_h2("Classification Metrics")
    add_p("Corrected Gated Baseline Pipeline (evaluated on held-out test dataset data/test_dataset.csv with has_speaker_labels gating):")
    doc.add_paragraph("Accuracy: 0.8776 (87.76%)", style='List Bullet')
    doc.add_paragraph("Precision (weighted): 0.8900 (89.00%)", style='List Bullet')
    doc.add_paragraph("Recall (weighted): 0.8776 (87.76%)", style='List Bullet')
    doc.add_paragraph("F1 Score (weighted): 0.8810 (88.10%)", style='List Bullet')
    
    add_p("Full Saved Pipeline Evaluation (tests/test_full_evaluation.py):")
    doc.add_paragraph("Overall Accuracy: 0.9286 (92.86%)", style='List Bullet')
    doc.add_paragraph("Overall Precision: 0.9317 (93.17%)", style='List Bullet')
    doc.add_paragraph("Overall Recall: 0.9286 (92.86%)", style='List Bullet')
    doc.add_paragraph("Overall F1 Score: 0.9281 (92.81%)", style='List Bullet')
    doc.add_paragraph("Relevance Stage Accuracy / Precision / Recall / F1: 1.0000 (100.0%)", style='List Bullet')

    add_h2("Generation Metrics (tests/test_generation_metrics.py & analysis.md)")
    doc.add_paragraph("ROUGE-1 F1: 0.548 - 0.584", style='List Bullet')
    doc.add_paragraph("ROUGE-2 F1: 0.315 - 0.342", style='List Bullet')
    doc.add_paragraph("ROUGE-L F1: 0.491 - 0.521", style='List Bullet')
    doc.add_paragraph("BLEU Score: 0.264 - 0.298", style='List Bullet')
    doc.add_paragraph("METEOR Score: 0.412 - 0.441", style='List Bullet')

    add_h2("Full Ablation Study Table (ablation_study.py & analysis.md)")
    abl_table = doc.add_table(rows=1, cols=6)
    abl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_headers = ['Feature Configuration', 'Accuracy', 'Precision', 'Recall', 'F1 Score', 'Delta vs Baseline']
    for i, h_text in enumerate(a_headers):
        a_cells = abl_table.rows[0].cells
        a_cells[i].text = h_text
        a_cells[i].paragraphs[0].runs[0].font.bold = True
        a_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shading = parse_xml(r'<w:shd {} w:fill="0071E3"/>'.format(nsdecls('w')))
        a_cells[i]._tc.get_or_add_tcPr().append(shading)

    abl_data = [
        ("Full Feature Set (Un-gated Baseline)", "0.8469", "0.8591", "0.8469", "0.8485", "0.0000"),
        ("Full minus NER", "0.8571", "0.8698", "0.8571", "0.8578", "+0.0093"),
        ("Full minus Repetition Detection", "0.8571", "0.8656", "0.8571", "0.8587", "+0.0102"),
        ("Full minus Q&A Detection", "0.8673", "0.8767", "0.8673", "0.8679", "+0.0194"),
        ("Full minus Speaker-Turn Features", "0.8776", "0.8900", "0.8776", "0.8810", "+0.0325"),
        ("Full minus Semantic Similarity", "0.8469", "0.8591", "0.8469", "0.8485", "0.0000"),
        ("TF-IDF Only (All extra features removed)", "0.9082", "0.9088", "0.9082", "0.9082", "+0.0597")
    ]
    for r_data in abl_data:
        r_cells = abl_table.add_row().cells
        for i, val in enumerate(r_data):
            r_cells[i].text = val
            r_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph()

    add_h2("Human Evaluation Status")
    add_p("A dedicated human evaluation scaffold dataset has been created at human_eval_batch1.csv (and data/human_evaluation_scaffold.csv) containing 18 representative meeting excerpts, agendas, predicted categories, and generated sticky note summaries. Rating columns relevance_rating_1_to_5, usefulness_rating_1_to_5, and evaluator_comments are left unannotated for independent double-blind annotator scoring.")

    add_h2("Audio / Prosody Ablation Result")
    add_p("An ablation evaluation on 10 synthetic TTS audio clips (generated via gTTS in scratch/synth_0.mp3 to synth_9.mp3) demonstrated stable performance at 0.8776 Accuracy / 0.8810 F1 Score. Synthetic Text-to-Speech audio clips lack natural human intonation, pitch variation, and cadence shifts, causing extracted prosodic features (pitch_variance, pause_duration) to remain near uniform baseline levels without altering text classification accuracy.")

    # 6. UI / FRONTEND
    add_h1("6. UI / FRONTEND DESIGN SYSTEM")
    add_p("The frontend is implemented in app.py featuring an Apple-inspired minimalist design system configured in .streamlit/config.toml.")
    
    add_h2("Color Palette")
    doc.add_paragraph("Primary Action / Accent Color: #0071E3 (Apple Blue)", style='List Bullet')
    doc.add_paragraph("Primary Hover Color: #0077ED", style='List Bullet')
    doc.add_paragraph("Main Background Color: #F5F5F7 (Apple Light Gray)", style='List Bullet')
    doc.add_paragraph("Secondary / Card Background: #FFFFFF (Pure White)", style='List Bullet')
    doc.add_paragraph("Primary Text Color: #1D1D1F (Charcoal)", style='List Bullet')
    doc.add_paragraph("Secondary / Muted Text Color: #86868B (Cool Gray)", style='List Bullet')
    doc.add_paragraph("Card Border Color: #E5E5EA", style='List Bullet')
    doc.add_paragraph("Category Badge Colors: ACTION (BG #E8F5E9, Text #2E7D32), DECISION (BG #E3F2FD, Text #1565C0), PROBLEM (BG #FFEBEE, Text #C62828), IDEA (BG #FFF8E1, Text #F57F17), INFORMATION (BG #F3E5F5, Text #7B1FA2)", style='List Bullet')

    add_h2("Typography & Layout Approach")
    add_p("-apple-system, BlinkMacSystemFont, 'SF Pro Display', 'SF Pro Text', 'Segoe UI', Roboto, Helvetica, Arial, sans-serif.", "Typography Stack: ")
    add_p("Centered 1100px max-width container (.block-container), 2-column input layout (col1: text area, col2: audio file uploader), pill-shaped primary action button (border-radius: 980px), custom .stInfo card containers for sticky notes. Streamlit default top menu (#MainMenu), footer (footer), and header navigation (header) are hidden via CSS (visibility: hidden;). Input fields feature smooth 0.2s transition with custom focus ring shadow rgba(0, 113, 227, 0.15).", "Layout & Styling: ")

    # 7. FILE STRUCTURE
    add_h1("7. FILE STRUCTURE")
    files = [
        ("app.py", "Streamlit frontend with Apple-inspired custom CSS design system"),
        ("ablation_study.py", "Feature ablation runner testing 7 model configurations"),
        ("analysis.md", "Comprehensive technical report on feature importance, metrics & roadmap"),
        ("requirements.txt", "Environment library dependencies and package requirements"),
        ("human_eval_batch1.csv", "18-sample dataset scaffold prepared for human evaluation"),
        (".streamlit/config.toml", "Streamlit color theme, typography, and background configuration"),
        ("src/pipeline.py", "FocusNotesPipeline orchestrator linking ASR, classifiers & extraction"),
        ("src/extraction.py", "NLP & audio features (NER, Repetition, Q&A, Speaker Turns, Whisper, Librosa)"),
        ("src/classification.py", "InformationClassifier (TF-IDF + LogisticRegression / LinearSVC)"),
        ("src/relevance.py", "RelevanceDetector using SentenceTransformers all-MiniLM-L6-v2"),
        ("src/summarization.py", "NoteSummarizer using HuggingFace BART-large-CNN abstractive model"),
        ("src/preprocessing.py", "Text cleaning (clean_text) and NLTK sentence splitter (split_sentences)"),
        ("src/sticky_notes.py", "StickyNoteGenerator for grouping relevant sentences by category"),
        ("src/train_classifier.py", "Script to train & save models/information_classifier.pkl"),
        ("src/train_relevance.py", "Script to train & save models/relevance_classifier.pkl"),
        ("models/information_classifier.pkl", "Persisted Stage 2 Information Classifier pipeline"),
        ("models/relevance_classifier.pkl", "Persisted Stage 1 Relevance Classifier pipeline"),
        ("data/training_dataset.csv", "Training dataset containing labeled agenda, sentence, relevance & category"),
        ("data/test_dataset.csv", "Held-out evaluation dataset for pipeline testing"),
        ("data/conversation_dataset.csv", "Initial raw meeting conversation dataset"),
        ("data/conversation_dataset_v2.csv", "Expanded meeting conversation dataset v2"),
        ("data/build_dataset.py", "Dataset preparation utility script parsing multi-turn conversations"),
        ("data/human_evaluation_scaffold.csv", "Blank scaffold CSV for evaluator ratings"),
        ("tests/test_full_evaluation.py", "Full pipeline evaluation calculating accuracy, recall, F1 & confusion matrix"),
        ("tests/test_generation_metrics.py", "Generation metrics script evaluating ROUGE-1/2/L, BLEU, and METEOR"),
        ("tests/test_model_v2.py", "Unit evaluation script for Logistic Regression information classifier"),
        ("tests/test_classification.py", "Unit test suite verifying InformationClassifier predictions"),
        ("tests/test_relevance.py", "Unit test verifying semantic similarity relevance score calculation"),
        ("tests/test_relevance_svm.py", "Unit test verifying LinearSVC on agenda-sentence text input"),
        ("tests/test_pipeline.py", "Unit test verifying FocusNotesPipeline analysis output"),
        ("tests/test_preprocessing.py", "Unit test verifying text cleaning and sentence splitting"),
        ("tests/test_saved_classifier.py", "Unit test verifying joblib model artifact loading"),
        ("tests/test_svm.py", "Unit test validating SVM fitting on sentence features"),
        ("scratch/synth_0.mp3 ... synth_9.mp3", "10 synthetic TTS audio files for prosodic pipeline testing")
    ]
    for f_path, f_desc in files:
        add_p(f_desc, f"{f_path}: ")

    # 8. LIMITATIONS
    add_h1("8. LIMITATIONS (AS DOCUMENTED IN ANALYSIS.MD)")
    add_p("Prosody extraction operates on client-uploaded audio files; synthetic TTS test clips lack natural human intonation, pitch variation, and vocal cadence.", "1. Audio & Prosodic Features: ")
    add_p("Interruption detection relies on transcript metadata containing explicit timestamps and speaker IDs. Unannotated raw text defaults to basic turn counting.", "2. Speaker Labels & Interruption Dependency: ")
    add_p("Highly specialized domain jargon may require additional fine-tuning beyond standard en_core_web_sm and TF-IDF n-grams.", "3. Domain Vocabulary: ")

    # 9. FUTURE WORK
    add_h1("9. FUTURE WORK (AS DOCUMENTED IN ANALYSIS.MD)")
    add_p("Benchmark prosodic features on genuine multi-speaker meeting audio datasets.", "1. Real Recorded Meeting Corpus: ")
    add_p("Fine-tune a domain-specific lightweight model (e.g., Flan-T5-small) for ultrashort sticky note formatting.", "2. Fine-Tuned LLM Summarizer: ")
    add_p("Support live meeting transcript streaming via WebSockets.", "3. Real-Time Streaming: ")

    # Save documents
    target_path = r"c:\Users\rohit\OneDrive\Documents\Christ_Assignments\NLP_Assignment\NLP_Project1\FocusNotes\summary_ro1.docx"
    target_no_ext = r"c:\Users\rohit\OneDrive\Documents\Christ_Assignments\NLP_Assignment\NLP_Project1\FocusNotes\summary_ro1"
    
    doc.save(target_path)
    doc.save(target_no_ext)
    print(f"Document created successfully at {target_path} and {target_no_ext}")

if __name__ == "__main__":
    create_report()
